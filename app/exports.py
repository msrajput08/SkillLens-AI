"""Export user-edited resume draft to a docx file (in-memory)."""
from io import BytesIO
from docx import Document
from docx.shared import Pt


def draft_docx(text: str) -> bytes:
    document = Document()
    sec = document.sections[0]
    sec.header_distance = Pt(28)
    for line in text.splitlines():
        clean = line.strip()
        if not clean:
            continue
        if clean.isupper() and 4 <= len(clean) <= 65 and not any(x in clean for x in ('@', ':')):
            document.add_heading(clean.title(), level=2)
        else:
            p = document.add_paragraph(clean)
            p.paragraph_format.space_after = Pt(4)
    output = BytesIO()
    document.save(output)
    return output.getvalue()
