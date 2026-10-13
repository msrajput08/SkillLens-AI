"""Safe-ish in-memory document parsing with limits; no file writes."""

from io import BytesIO
from zipfile import BadZipFile, ZipFile
from pypdf import PdfReader
from docx import Document

MAX_UPLOAD = 6 * 1024 * 1024
MAX_PDF_PAGES = 20
MAX_DOCX_UNCOMPRESSED = 35 * 1024 * 1024

class ResumeParseError(ValueError):
    pass


def extract_resume(filename: str, content: bytes) -> str:
    suffix = filename.rsplit('.', 1)[-1].lower() if '.' in filename else ''
    if not content:
        raise ResumeParseError("The uploaded file is empty.")
    if len(content) > MAX_UPLOAD:
        raise ResumeParseError("Resume must be 6 MB or smaller.")
    if suffix == 'pdf':
        if not content.startswith(b'%PDF-'):
            raise ResumeParseError("This file does not appear to be a valid PDF.")
        try:
            pdf = PdfReader(BytesIO(content), strict=False)
            if pdf.is_encrypted:
                raise ResumeParseError("Password-protected PDFs are not supported.")
            if len(pdf.pages) > MAX_PDF_PAGES:
                raise ResumeParseError("PDF must be 20 pages or fewer.")
            text = '\n'.join(page.extract_text(extraction_mode='plain') or '' for page in pdf.pages)
        except ResumeParseError:
            raise
        except Exception:
            raise ResumeParseError("Could not extract this PDF. Try a text-based PDF or paste the resume text.")
    elif suffix == 'docx':
        try:
            with ZipFile(BytesIO(content)) as archive:
                if len(archive.infolist()) > 1200 or sum(i.file_size for i in archive.infolist()) > MAX_DOCX_UNCOMPRESSED:
                    raise ResumeParseError("DOCX file is too complex or large to process.")
                if 'word/document.xml' not in archive.namelist():
                    raise ResumeParseError("Not a valid Word document.")
            doc = Document(BytesIO(content))
            lines = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    lines.append(' | '.join(cell.text for cell in row.cells))
            text = '\n'.join(lines)
        except ResumeParseError:
            raise
        except (BadZipFile, Exception):
            raise ResumeParseError("Could not read this DOCX file. Try pasting your resume text.")
    elif suffix == 'txt':
        try:
            text = content.decode('utf-8-sig')
        except UnicodeDecodeError:
            raise ResumeParseError("TXT files must use UTF-8 encoding.")
    else:
        raise ResumeParseError("Supported formats: PDF, DOCX and TXT.")
    if len(text.strip()) < 100:
        raise ResumeParseError("Not enough readable text found. Scanned/image-only PDFs need OCR; paste resume text instead.")
    return text[:55000]
